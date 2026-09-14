from pathlib import Path
import time

import pandas as pd
from openpyxl import Workbook

from cost.cost_database_csv import (
    csv_path_for_sheet,
    read_cost_database_sheet,
    sync_cost_database_csvs,
)


def _write_workbook(path: Path, unit_cost: float = 12.5) -> None:
    workbook = Workbook()
    costs = workbook.active
    costs.title = "Cost Database"
    costs.append(["Account", "Unit Cost", "Description"])
    costs.append([10, unit_cost, "Fuel, handling"])

    inflation = workbook.create_sheet("Inflation Adjustment")
    inflation.append(["Year", "General"])
    inflation.append([2025, 1.0])

    economics = workbook.create_sheet("Economics Parameters")
    economics.append(["Parameter", "Value"])
    economics.append(["Interest Rate", 0.05])
    workbook.save(path)


def test_sync_exports_every_sheet_and_reads_primary_csv(tmp_path):
    workbook_path = tmp_path / "Cost_Database.xlsx"
    _write_workbook(workbook_path)

    paths = sync_cost_database_csvs(workbook_path)

    expected_directory = tmp_path.resolve()
    assert paths["Cost Database"] == expected_directory / "Cost_Database.csv"
    assert paths["Inflation Adjustment"] == (
        expected_directory / "Cost_Database__Inflation_Adjustment.csv"
    )
    assert set(paths) == {
        "Cost Database",
        "Inflation Adjustment",
        "Economics Parameters",
    }
    frame = read_cost_database_sheet(workbook_path, "Cost Database")
    assert frame.loc[0, "Unit Cost"] == 12.5
    assert frame.loc[0, "Description"] == "Fuel, handling"


def test_sync_updates_changed_values_without_touching_unchanged_csv(tmp_path):
    workbook_path = tmp_path / "Cost_Database.xlsx"
    _write_workbook(workbook_path)
    paths = sync_cost_database_csvs(workbook_path)
    primary_csv = paths["Cost Database"]
    first_mtime = primary_csv.stat().st_mtime_ns

    # An unchanged sync must not dirty the file or alter its timestamp.
    sync_cost_database_csvs(workbook_path)
    assert primary_csv.stat().st_mtime_ns == first_mtime

    time.sleep(0.002)
    _write_workbook(workbook_path, unit_cost=14.75)
    updated = read_cost_database_sheet(workbook_path, "Cost Database")
    assert updated.loc[0, "Unit Cost"] == 14.75
    assert primary_csv.stat().st_mtime_ns > first_mtime


def test_csv_paths_are_stable_and_sheet_specific(tmp_path):
    workbook_path = tmp_path / "Any Name.xlsx"
    assert csv_path_for_sheet(workbook_path, "Cost Database") == (
        tmp_path / "Any Name.csv"
    )
    assert csv_path_for_sheet(workbook_path, "Economics Parameters") == (
        tmp_path / "Any Name__Economics_Parameters.csv"
    )
