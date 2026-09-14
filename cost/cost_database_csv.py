# Copyright 2025, Battelle Energy Alliance, LLC, ALL RIGHTS RESERVED
"""Deterministic CSV mirror for the Excel cost database.

Excel remains the human-edited source of truth.  Before cost data is read,
this module exports every workbook sheet to a stable CSV path so that changes
are reviewable in Git and the cost engine consumes the same values that were
exported.
"""

from __future__ import annotations

import argparse
import csv
from datetime import date, datetime, time
import io
import os
from pathlib import Path
import re
import tempfile
import threading
from typing import Dict

import pandas as pd
from openpyxl import load_workbook


_PRIMARY_SHEET = "Cost Database"
_sync_lock = threading.Lock()
_last_sync: dict[Path, tuple[tuple[int, int], Dict[str, Path]]] = {}


def _sheet_suffix(sheet_name: str) -> str:
    """Return a filesystem-safe, stable suffix for a worksheet name."""
    suffix = re.sub(r"[^A-Za-z0-9]+", "_", sheet_name).strip("_")
    if not suffix:
        raise ValueError(f"Worksheet name {sheet_name!r} has no usable characters")
    return suffix


def csv_path_for_sheet(workbook_path: str | os.PathLike[str], sheet_name: str) -> Path:
    """Return the CSV mirror path for ``sheet_name``."""
    workbook = Path(workbook_path)
    if sheet_name == _PRIMARY_SHEET:
        return workbook.with_suffix(".csv")
    return workbook.with_name(f"{workbook.stem}__{_sheet_suffix(sheet_name)}.csv")


def _csv_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    return value


def _worksheet_csv_bytes(worksheet) -> bytes:
    """Serialize cached worksheet values with deterministic CSV settings."""
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]

    # openpyxl's worksheet bounds include cells that have formatting but no
    # value. pandas.read_excel historically discarded those trailing rows and
    # columns, so trim them to preserve the cost engine's existing schema.
    while rows and all(value is None for value in rows[-1]):
        rows.pop()
    last_value_column = max(
        (
            index
            for row in rows
            for index, value in enumerate(row)
            if value is not None
        ),
        default=-1,
    )

    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    for row in rows:
        writer.writerow(
            [_csv_value(value) for value in row[: last_value_column + 1]]
        )
    return output.getvalue().encode("utf-8")


def _replace_if_changed(path: Path, content: bytes) -> bool:
    if path.exists() and path.read_bytes() == content:
        return False

    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "wb") as temporary_file:
            temporary_file.write(content)
            temporary_file.flush()
            os.fsync(temporary_file.fileno())
        os.replace(temporary_path, path)
    except Exception:
        temporary_path.unlink(missing_ok=True)
        raise
    return True


def sync_cost_database_csvs(
    workbook_path: str | os.PathLike[str],
) -> Dict[str, Path]:
    """Regenerate CSV mirrors from an Excel cost database when it changes.

    Formulas are exported as the values cached by Excel, matching the values
    pandas/openpyxl historically returned to the cost engine. Files are
    replaced atomically and are left untouched when their content is identical.
    """
    workbook = Path(workbook_path).resolve()
    if workbook.suffix.lower() != ".xlsx":
        raise ValueError(f"Cost database must be an .xlsx workbook: {workbook}")
    if not workbook.is_file():
        raise FileNotFoundError(f"Cost database workbook not found: {workbook}")

    stat = workbook.stat()
    signature = (stat.st_mtime_ns, stat.st_size)

    with _sync_lock:
        cached = _last_sync.get(workbook)
        if cached is not None and cached[0] == signature:
            # Ensure a generated file was not removed since the previous call
            # in this process, without reopening the workbook on every read.
            paths = cached[1]
            if all(path.is_file() for path in paths.values()):
                return paths

        excel = load_workbook(workbook, read_only=True, data_only=True)
        try:
            if _PRIMARY_SHEET not in excel.sheetnames:
                raise ValueError(
                    f"Cost database is missing required sheet {_PRIMARY_SHEET!r}"
                )
            rendered = {
                sheet_name: (
                    csv_path_for_sheet(workbook, sheet_name),
                    _worksheet_csv_bytes(excel[sheet_name]),
                )
                for sheet_name in excel.sheetnames
            }
        finally:
            excel.close()

        for path, content in rendered.values():
            _replace_if_changed(path, content)

        paths = {name: path for name, (path, _) in rendered.items()}
        _last_sync[workbook] = (signature, paths)
        return paths


def read_cost_database_sheet(
    workbook_path: str | os.PathLike[str], sheet_name: str
) -> pd.DataFrame:
    """Sync the workbook and read one sheet from its generated CSV mirror."""
    csv_paths = sync_cost_database_csvs(workbook_path)
    try:
        csv_path = csv_paths[sheet_name]
    except KeyError as exc:
        raise ValueError(f"Cost database has no sheet named {sheet_name!r}") from exc
    return pd.read_csv(csv_path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Regenerate deterministic CSV mirrors of a cost workbook."
    )
    parser.add_argument(
        "workbook",
        nargs="?",
        default="cost/Cost_Database.xlsx",
        help="cost workbook to synchronize (default: cost/Cost_Database.xlsx)",
    )
    args = parser.parse_args()
    paths = sync_cost_database_csvs(args.workbook)
    for sheet_name, path in paths.items():
        print(f"{sheet_name}: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
